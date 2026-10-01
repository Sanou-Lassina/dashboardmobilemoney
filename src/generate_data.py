"""Générateur de données SIMULÉES mais plausibles pour un opérateur Mobile Money au Burkina Faso.

Ce qui rend la simulation crédible (et testable) :
- montants log-normaux (beaucoup de petits montants, peu de gros) et paliers de frais ;
- saisonnalité : jour de semaine, fin de mois (salaires), Ramadan/Tabaski, rentrée scolaire ;
- profil horaire (pics matin et fin d'après-midi) ;
- clients très inégaux (loi de Pareto), inscriptions progressives, départs (churn) ;
- réseau d'agents avec agents dormants, déséquilibres cash-in / cash-out par zone ;
- échecs motivés (solde, PIN, réseau, plafond, float) et 4 pannes datées ;
- anomalies injectées (fractionnement sous plafond, rafales, gros montants nocturnes) ;
- défauts de qualité volontaires (valeurs manquantes, doublons, orphelins) pour le DQA.

Usage :  python -m src.generate_data [--sortie data/Base_MobileMoney.xlsx] [--graine 42]
"""
import argparse
from pathlib import Path

import numpy as np
import pandas as pd

from . import config as C

DEBUT, FIN = pd.Timestamp("2024-01-01"), pd.Timestamp("2025-12-31")

# nom, ville, région, lat, lon, poids, biais cash-out, performance relative, code
AGENCES = [
    ("Ouaga Centre", "Ouagadougou", "Centre", 12.3714, -1.5197, 0.22, -0.15, 1.00, "OUC"),
    ("Ouaga 2000", "Ouagadougou", "Centre", 12.3010, -1.4900, 0.16, -0.10, 1.05, "O2K"),
    ("Bobo-Dioulasso", "Bobo-Dioulasso", "Hauts-Bassins", 11.1771, -4.2979, 0.18, 0.00, 0.98, "BBD"),
    ("Koudougou", "Koudougou", "Centre-Ouest", 12.2526, -2.3627, 0.10, 0.10, 0.95, "KDG"),
    ("Ouahigouya", "Ouahigouya", "Nord", 13.5828, -2.4216, 0.09, 0.20, 0.88, "OHG"),
    ("Banfora", "Banfora", "Cascades", 10.6333, -4.7667, 0.10, 0.10, 1.02, "BFA"),
    ("Fada N'Gourma", "Fada N'Gourma", "Est", 12.0616, 0.3583, 0.09, 0.25, 0.90, "FDA"),
    ("Dori", "Dori", "Sahel", 14.0354, -0.0346, 0.06, 0.30, 0.80, "DOR"),
]
NOMS = ["Ouédraogo", "Traoré", "Sawadogo", "Kaboré", "Compaoré", "Zongo", "Sanou", "Ouattara", "Nikiéma", "Tapsoba",
        "Bationo", "Somé", "Barro", "Diallo", "Sankara", "Konaté", "Coulibaly", "Zoungrana", "Belem", "Yaméogo",
        "Kinda", "Dabiré", "Sorgho", "Lankoandé", "Ilboudo", "Savadogo", "Bambara", "Tiendrébéogo", "Hien", "Da"]
PRENOMS_M = ["Issa", "Moussa", "Adama", "Salif", "Drissa", "Ibrahim", "Désiré", "Honoré", "Seydou", "Boureima",
             "Abdoulaye", "Souleymane", "Alain", "Eric", "Jean", "Paul", "Idrissa", "Mahamadi", "Harouna", "Lassané"]
PRENOMS_F = ["Aminata", "Fatimata", "Awa", "Blandine", "Rasmata", "Mariam", "Judith", "Salamata", "Estelle", "Alimata",
             "Nafissatou", "Pauline", "Odile", "Habibou", "Safiatou", "Aïcha", "Clarisse", "Nadège", "Rokia", "Zénabou"]
SEGMENTS = ["Grand public", "Commerçant", "Salarié", "Rural / Agriculteur", "Étudiant"]
P_SEG = [0.42, 0.16, 0.18, 0.14, 0.10]
EVENEMENTS = [pd.Timestamp(d) for d in ("2024-04-10", "2024-06-16", "2025-03-30", "2025-06-06")]  # Eid / Tabaski
# (date de début, type, périmètre, durée en jours)
PANNES = [("2024-05-14", "canal", ["USSD"], 3), ("2024-11-03", "agence", ["Bobo-Dioulasso"], 2),
          ("2025-02-20", "canal", ["App Mobile"], 2), ("2025-08-12", "agence", ["Ouaga Centre"], 2)]


def _commissions(prod, canal, m, ok):
    """Frais clients (revenu brut) et commission versée à l'agent, selon des paliers réglementés plausibles."""
    m = np.asarray(m, float)
    frais = np.select(
        [prod == C.P_DEPOT, prod == C.P_RETRAIT, prod == C.P_P2P, prod == C.P_FACTURE, prod == C.P_CREDIT,
         prod == C.P_MARCHAND],
        [0.0, np.clip(0.012 * m, 50, 7500), np.clip(0.010 * m, 25, 3000), np.clip(0.008 * m, 25, 2000),
         0.03 * m, np.clip(0.009 * m, 10, 5000)], 0.0)
    frais = np.where(ok, np.round(frais), 0.0)
    via_agent = canal == C.CANAL_AGENT
    com_ag = np.select([(prod == C.P_DEPOT) & ok & via_agent, (prod == C.P_RETRAIT) & ok & via_agent, via_agent & ok],
                       [0.003 * m, 0.40 * frais, 0.25 * frais], 0.0)
    return frais, np.round(com_ag)


def _arrondir(prod, m):
    pas = np.where((prod == C.P_DEPOT) | (prod == C.P_RETRAIT), 500, 50)
    return np.clip(np.maximum(np.round(m / pas) * pas, pas), 100, 990_000)


def generer(chemin, graine: int = 42, n_tx: int = 50_000, n_clients: int = 2_200) -> Path:
    chemin = Path(chemin)
    rng = np.random.default_rng(graine)
    na = len(AGENCES)
    poids_ag = np.array([a[5] for a in AGENCES])
    nom_ag = np.array([a[0] for a in AGENCES])

    # ───────────── Réseau : commerciaux et agents ─────────────
    reps_pool = [f"{rng.choice(NOMS)} {rng.choice(PRENOMS_M + PRENOMS_F)}" for _ in range(40)]
    reps_pool = list(dict.fromkeys(reps_pool))
    while len(reps_pool) < na * 4:
        reps_pool.append(f"{rng.choice(NOMS)} {rng.choice(PRENOMS_M + PRENOMS_F)}")
    reps = {i: reps_pool[4 * i:4 * i + 4] for i in range(na)}
    ag_rows = []
    for i, a in enumerate(AGENCES):
        for k in range(int(round(a[5] * 90)) + 4):
            tot_av = rng.random() < 0.85
            act = (pd.Timestamp("2019-01-01") + pd.Timedelta(days=int(rng.integers(0, 365 * 5)))) if tot_av else \
                  (pd.Timestamp("2024-01-15") + pd.Timedelta(days=int(rng.integers(0, 600))))
            ag_rows.append({
                "ID_Agent": f"AG-{a[8]}-{k + 1:03d}", "Nom_Point": f"Point Mobile Money {a[1]} {k + 1}",
                "Agence": a[0], "Ville": a[1], "Commercial": reps[i][k % 4], "Date_Activation": act,
                "Latitude": a[3] + rng.normal(0, 0.03), "Longitude": a[4] + rng.normal(0, 0.03),
                "_i": i, "_w": 0.0 if rng.random() < 0.08 else rng.lognormal(0, 0.9)})
    agents = pd.DataFrame(ag_rows)

    # ───────────── Clients ─────────────
    seg = rng.choice(SEGMENTS, n_clients, p=P_SEG)
    sexe = np.where(rng.random(n_clients) < np.where(seg == "Commerçant", 0.35, 0.44), "F", "M")
    age = np.clip(np.where(seg == "Étudiant", rng.normal(22, 3, n_clients), rng.normal(35, 11, n_clients)), 18, 72).round()
    ag_cl = rng.choice(na, n_clients, p=poids_ag)
    avant = rng.random(n_clients) < 0.60
    ins = np.where(avant,
                   pd.Timestamp("2021-06-01").value + (rng.random(n_clients) * (pd.Timestamp("2023-12-31").value - pd.Timestamp("2021-06-01").value)),
                   DEBUT.value + (rng.random(n_clients) * ((FIN - pd.Timedelta(days=10)).value - DEBUT.value)))
    ins = pd.to_datetime(ins.astype("int64")).normalize()
    quitte = rng.random(n_clients) < 0.30
    fin_act = pd.to_datetime(np.where(quitte, (ins + pd.to_timedelta(rng.integers(60, 700, n_clients), unit="D")).values,
                                      np.datetime64("2030-01-01")))
    mult_seg = pd.Series(seg).map({"Commerçant": 3.0, "Salarié": 1.3, "Grand public": 1.0,
                                   "Rural / Agriculteur": 0.6, "Étudiant": 0.7}).values
    w_cl = rng.lognormal(0, 0.85, n_clients) * mult_seg
    w_cl[rng.random(n_clients) < 0.05] = 0          # inscrits jamais actifs
    kyc = np.where(seg == "Commerçant", rng.choice(["Niveau 1", "Niveau 2", "Niveau 3"], n_clients, p=[.2, .5, .3]),
                   rng.choice(["Niveau 1", "Niveau 2", "Niveau 3"], n_clients, p=[.6, .32, .08]))
    clients = pd.DataFrame({
        "ID_Client": [f"CLI-{i + 1:06d}" for i in range(n_clients)],
        "Nom": rng.choice(NOMS, n_clients),
        "Prénom": np.where(sexe == "F", rng.choice(PRENOMS_F, n_clients), rng.choice(PRENOMS_M, n_clients)),
        "Sexe": sexe, "Age": age, "Segment": seg, "KYC_Niveau": kyc, "Date_Inscription": ins,
        "Agence": nom_ag[ag_cl], "Ville": [AGENCES[i][1] for i in ag_cl],
        "Commercial": [reps[i][int(rng.integers(0, 4))] for i in ag_cl]})

    # ───────────── Calendrier de l'activité ─────────────
    jours = pd.date_range(DEBUT, FIN, freq="D")
    mi_j = np.asarray((jours.year - 2024) * 12 + jours.month - 1)
    w_jour = np.array([1.0, 0.95, 0.95, 1.0, 1.15, 1.2, 0.7])[np.asarray(jours.dayofweek)]
    w_jour = w_jour * np.where((np.asarray(jours.day) >= 25) | (np.asarray(jours.day) <= 3), 1.18, 1.0) * (1 + 0.02 * mi_j)
    for e in EVENEMENTS:
        w_jour = w_jour * np.where(np.asarray((jours >= e - pd.Timedelta(days=5)) & (jours <= e)), 1.5, 1.0)
    comptes = rng.multinomial(n_tx, w_jour / w_jour.sum())
    jour_tx = np.repeat(jours.values, comptes)
    N = len(jour_tx)
    h = np.array([.2, .1, .05, .05, .1, .4, 1, 2, 3.5, 4, 3.8, 3.2, 3, 3.2, 3, 3, 3.5, 4, 3.8, 3, 2, 1.2, .7, .4])
    heure = rng.choice(24, N, p=h / h.sum())
    ts = pd.Series(pd.to_datetime(jour_tx) + pd.to_timedelta(heure, unit="h") +
                   pd.to_timedelta(rng.integers(0, 60, N), unit="m") + pd.to_timedelta(rng.integers(0, 60, N), unit="s"))

    # ───────────── Client de chaque transaction (inscrit, pas encore parti, pondéré par l'activité) ─────────────
    mkey = (ts.dt.year * 12 + ts.dt.month).values
    cid = np.empty(N, int)
    for mk in np.unique(mkey):
        idx = np.where(mkey == mk)[0]
        m0 = pd.Timestamp(year=int((mk - 1) // 12), month=int((mk - 1) % 12 + 1), day=1)
        m1 = m0 + pd.offsets.MonthEnd(0)
        elig = np.where((ins <= m1) & (fin_act >= m0) & (w_cl > 0))[0]
        p = w_cl[elig] / w_cl[elig].sum()
        cid[idx] = rng.choice(elig, len(idx), p=p)
    ins_tx = pd.Series(ins[cid])
    tod = ts - ts.dt.normalize()
    ts = ts.where(ts >= ins_tx, ins_tx.dt.normalize() + tod)       # pas de transaction avant l'inscription
    mi = ((ts.dt.year - 2024) * 12 + ts.dt.month - 1).values

    seg_tx = seg[cid]
    ag_tx = ag_cl[cid].copy()
    roam = rng.random(N) < 0.12
    ag_tx[roam] = rng.choice(na, roam.sum(), p=poids_ag)
    biais = np.array([a[6] for a in AGENCES])[ag_tx]

    # ───────────── Produit ─────────────
    P = np.tile(np.array([.17, .21, .24, .10, .17, .11]), (N, 1))
    mult = {"Commerçant": [1.5, 1, .8, .7, .6, 2.5], "Salarié": [1, 1, 1.1, 1.6, 1, 1],
            "Rural / Agriculteur": [.9, 1.5, .9, .6, .9, .3], "Étudiant": [.8, .7, 1.2, .4, 1.6, .8]}
    for s_, m_ in mult.items():
        P[seg_tx == s_] *= np.array(m_)
    P[:, 1] *= 1 + 2 * biais
    P[:, 0] *= np.clip(1 - 1.2 * biais, .3, None)
    P[:, 5] *= 1 + 0.04 * mi
    P[:, 2] *= 1 + 0.01 * mi
    eid = np.zeros(N, bool)
    for e in EVENEMENTS:
        eid |= ((ts >= e - pd.Timedelta(days=5)) & (ts <= e + pd.Timedelta(days=1))).values
    P[eid, 2] *= 2.0
    rentree = (((ts.dt.month == 9) & (ts.dt.day >= 15)) | ((ts.dt.month == 10) & (ts.dt.day <= 15))).values
    P[rentree, 3] *= 1.8
    P /= P.sum(axis=1, keepdims=True)
    pi = (rng.random(N)[:, None] > P.cumsum(axis=1)).sum(axis=1).clip(0, 5)
    produit = np.array(C.PRODUITS, dtype=object)[pi]

    # ───────────── Canal ─────────────
    CP = np.array([[.97, .02, .01, 0], [.98, .01, .01, 0], [.15, .50, .35, 0],
                   [.25, .40, .35, 0], [.15, .65, .20, 0], [0, .10, .40, .50]])
    pc = CP[pi].copy()
    pc[:, 2] *= 1 + 0.03 * mi
    pc /= pc.sum(axis=1, keepdims=True)
    ci = (rng.random(N)[:, None] > pc.cumsum(axis=1)).sum(axis=1).clip(0, 3)
    canal = np.array(C.CANAUX, dtype=object)[ci]

    # ───────────── Agent (canal Agent) et commercial ─────────────
    id_agent = np.full(N, None, dtype=object)
    commercial = np.empty(N, dtype=object)
    mois_obs = ts.dt.year.values * 12 + ts.dt.month.values
    for i in range(na):
        ag_i = agents[agents["_i"] == i].reset_index(drop=True)
        for mk in np.unique(mkey):
            idx = np.where((ag_tx == i) & (canal == C.CANAL_AGENT) & (mkey == mk))[0]
            if not len(idx):
                continue
            m1 = pd.Timestamp(year=int((mk - 1) // 12), month=int((mk - 1) % 12 + 1), day=1) + pd.offsets.MonthEnd(0)
            el = ag_i[(ag_i["Date_Activation"] <= m1) & (ag_i["_w"] > 0)]
            if el.empty:
                el = ag_i[ag_i["_w"] > 0]
            ch = rng.choice(len(el), len(idx), p=(el["_w"] / el["_w"].sum()).values)
            id_agent[idx] = el["ID_Agent"].values[ch]
            commercial[idx] = el["Commercial"].values[ch]
    pas_agent = id_agent == None  # noqa: E711
    for i in range(na):
        idx = np.where(pas_agent & (ag_tx == i))[0]
        commercial[idx] = np.array(reps[i], dtype=object)[rng.integers(0, 4, len(idx))]

    # ───────────── Montants ─────────────
    med = np.select([produit == C.P_DEPOT, produit == C.P_RETRAIT, produit == C.P_P2P, produit == C.P_FACTURE],
                    [25_000, 20_000, 10_000, 15_000], 6_000).astype(float)
    sig = np.select([produit == C.P_DEPOT, produit == C.P_RETRAIT, produit == C.P_P2P, produit == C.P_FACTURE],
                    [1.0, 0.9, 1.0, 0.7], 0.9)
    ms = pd.Series(seg_tx).map({"Commerçant": 2.2, "Salarié": 1.3, "Rural / Agriculteur": 0.8, "Étudiant": 0.5}).fillna(1.0).values
    montant = np.exp(np.log(med * ms * (1 + 0.006 * mi)) + sig * rng.standard_normal(N))
    montant = _arrondir(produit, montant)
    cr = produit == C.P_CREDIT
    montant[cr] = rng.choice([100, 200, 500, 1000, 2000, 5000], cr.sum(), p=[.05, .10, .30, .30, .15, .10])

    # ───────────── Échecs : base, canal, heure de pointe, gros montants, pannes ─────────────
    d_norm = ts.dt.normalize()
    hh = ts.dt.hour.values
    mult_panne = np.ones(N)
    en_panne = np.zeros(N, bool)
    for dte, typ, val, duree in PANNES:
        jours_p = pd.Timestamp(dte) + pd.to_timedelta(np.arange(duree), unit="D")
        m_ = d_norm.isin(jours_p).values & (np.isin(canal, val) if typ == "canal" else np.isin(nom_ag[ag_tx], val))
        mult_panne[m_] = 15.0
        en_panne |= m_
    mc = pd.Series(canal).map({"USSD": 1.7, "App Mobile": 1.0, "Agent": 0.8, "API Marchand": 0.6}).values
    pointe = np.where(((hh >= 8) & (hh <= 10)) | ((hh >= 17) & (hh <= 19)), 1.3, 1.0)
    gros = np.where(montant > 200_000, 1.5, 1.0)
    p_ko = np.clip(0.035 * mc * pointe * gros * mult_panne, 0, 0.95)
    p_att = 0.015 * np.where(en_panne, 4, 1)
    u = rng.random(N)
    statut = np.where(u < p_ko, C.STATUT_KO, np.where(u < p_ko + p_att, C.STATUT_ATT, C.STATUT_OK)).astype(object)
    motif = np.full(N, None, dtype=object)
    ko = np.where(statut == C.STATUT_KO)[0]
    base = np.tile(np.array([.32, .18, .22, .08, .08, .12]), (len(ko), 1))
    is_ag = canal[ko] == C.CANAL_AGENT
    base[~is_ag, 5] = 0
    base[is_ag, 2] *= 0.5
    base[montant[ko] >= 300_000, 3] *= 3
    base[en_panne[ko], 2] *= 12
    base[produit[ko] == C.P_RETRAIT, 0] *= 1.5
    base[produit[ko] == C.P_DEPOT, 0] *= 0.1
    base /= base.sum(axis=1, keepdims=True)
    motif[ko] = np.array(C.MOTIFS_ECHEC, dtype=object)[(rng.random(len(ko))[:, None] > base.cumsum(axis=1)).sum(axis=1).clip(0, 5)]
    att = np.where(statut == C.STATUT_ATT)[0]
    motif[att] = np.where(rng.random(len(att)) < 0.8, "Timeout réseau", "Confirmation en cours")

    frais, com_ag = _commissions(produit, canal, montant, statut == C.STATUT_OK)
    tx = pd.DataFrame({
        "Date": ts, "ID_Client": clients["ID_Client"].values[cid], "Agence": nom_ag[ag_tx],
        "Ville": [AGENCES[i][1] for i in ag_tx], "Commercial": commercial, "ID_Agent": id_agent,
        "Produit": produit, "Canal": canal, "Montant (FCFA)": montant, "Commission (FCFA)": frais,
        "Commission_Agent (FCFA)": com_ag, "Statut": statut, "Motif_Echec": motif})

    # ───────────── Anomalies injectées (fraude / conformité) ─────────────
    extra = []
    def _extra(c_idx, t, prod, can, m, stat="Succès"):
        i = ag_cl[c_idx]
        ag_i = agents[(agents["_i"] == i) & (agents["_w"] > 0)]
        a_id = rng.choice(ag_i["ID_Agent"].values) if can == C.CANAL_AGENT else None
        extra.append({"Date": t, "ID_Client": clients["ID_Client"].iat[c_idx], "Agence": nom_ag[i],
                      "Ville": AGENCES[i][1], "Commercial": reps[i][0], "ID_Agent": a_id, "Produit": prod,
                      "Canal": can, "Montant (FCFA)": float(m), "Statut": stat, "Motif_Echec": None})
    candidats = np.where((w_cl > 0) & (ins < pd.Timestamp("2025-06-01")) & (~quitte))[0]
    for c_idx in rng.choice(candidats, 22, replace=False):                       # fractionnement sous plafond
        ancre = DEBUT + pd.Timedelta(days=int(rng.integers(60, 680)))
        for _ in range(int(rng.integers(4, 7))):
            t = ancre + pd.Timedelta(days=int(rng.integers(0, 6)), hours=int(rng.integers(9, 18)), minutes=int(rng.integers(0, 60)))
            retrait = rng.random() < 0.6
            _extra(c_idx, t, C.P_RETRAIT if retrait else C.P_P2P, C.CANAL_AGENT if retrait else "USSD",
                   round(rng.uniform(900_000, 995_000), -3))
    for c_idx in rng.choice(candidats, 12, replace=False):                       # rafales
        ancre = DEBUT + pd.Timedelta(days=int(rng.integers(30, 700)), hours=int(rng.integers(8, 20)))
        for _ in range(int(rng.integers(9, 15))):
            _extra(c_idx, ancre + pd.Timedelta(minutes=int(rng.integers(0, 40))), C.P_P2P, "USSD",
                   round(rng.uniform(1_000, 5_000), -2), "Succès" if rng.random() < 0.7 else "Échoué")
    for c_idx in rng.choice(candidats, 10, replace=False):                       # gros montants de nuit
        for _ in range(int(rng.integers(2, 5))):
            t = DEBUT + pd.Timedelta(days=int(rng.integers(30, 700)), hours=int(rng.integers(1, 5)), minutes=int(rng.integers(0, 60)))
            retrait = rng.random() < 0.5
            _extra(c_idx, t, C.P_RETRAIT if retrait else C.P_P2P, C.CANAL_AGENT if retrait else "USSD",
                   round(rng.uniform(250_000, 800_000), -3))
    ex = pd.DataFrame(extra)
    ex["Motif_Echec"] = np.where(ex["Statut"] == C.STATUT_KO, "PIN erroné", None)
    f_, a_ = _commissions(ex["Produit"].values, ex["Canal"].values, ex["Montant (FCFA)"].values, (ex["Statut"] == C.STATUT_OK).values)
    ex["Commission (FCFA)"], ex["Commission_Agent (FCFA)"] = f_, a_
    tx = pd.concat([tx, ex], ignore_index=True).sort_values("Date").reset_index(drop=True)
    tx.insert(1, "ID_Transaction", [f"TX-{i + 1:07d}" for i in range(len(tx))])

    # ───────────── Objectifs : agence × année × mois (réalisés × facteur de performance) ─────────────
    ok = tx[tx["Statut"] == C.STATUT_OK].copy()
    ok["Année"], ok["M"] = ok["Date"].dt.year, ok["Date"].dt.month
    reel = ok.groupby(["Agence", "Année", "M"]).agg(CA=("Montant (FCFA)", "sum"), Tx=("ID_Transaction", "count")).reset_index()
    tr_ag = (ok.groupby("Agence")["Commission (FCFA)"].sum() / ok.groupby("Agence")["Montant (FCFA)"].sum() * 100)
    perf = {a[0]: a[7] for a in AGENCES}
    reg = {a[0]: a[2] for a in AGENCES}
    obj_rows = []
    for _, r in reel.iterrows():
        f = perf[r["Agence"]] * np.exp(rng.normal(0, 0.07))
        ft = perf[r["Agence"]] * np.exp(rng.normal(0, 0.05))
        obj_rows.append({"Agence": r["Agence"], "Région": reg[r["Agence"]], "Année": int(r["Année"]),
                         "Mois": C.MOIS_NOMS[int(r["M"]) - 1], "N° Mois": int(r["M"]),
                         "Objectif_Montant (FCFA)": round(r["CA"] / f, -5), "Objectif_Transactions": int(round(r["Tx"] / ft, -1)),
                         "Taux_Commission_Cible (%)": round(tr_ag[r["Agence"]] * 1.05, 2)})
    objectifs = pd.DataFrame(obj_rows).sort_values(["Année", "N° Mois", "Agence"]).reset_index(drop=True)

    # ───────────── Défauts de qualité volontaires (alimentent le DQA) ─────────────
    n = len(tx)
    pick = lambda k: rng.choice(n, k, replace=False)  # noqa: E731
    tx.loc[pick(60), "Montant (FCFA)"] = np.nan
    tx.loc[pick(12), "Montant (FCFA)"] = -1000
    tx.loc[pick(10), "Montant (FCFA)"] = 0
    tx["Statut"] = tx["Statut"].astype(object)
    tx.loc[pick(120), "Statut"] = np.nan
    tx.loc[pick(200), "Ville"] = np.nan
    tx.loc[pick(15), "ID_Client"] = [f"CLI-9{int(x):05d}" for x in rng.integers(0, 99999, 15)]
    ic = pick(10)
    tx.loc[ic, "Commission (FCFA)"] = (tx.loc[ic, "Montant (FCFA)"].abs().fillna(1000) * 1.2).round()
    ech = tx.index[tx["Statut"] == C.STATUT_KO]
    tx.loc[rng.choice(ech, int(len(ech) * 0.06), replace=False), "Motif_Echec"] = np.nan
    tx = pd.concat([tx, tx.sample(25, random_state=graine)], ignore_index=True)   # doublons d'ID
    tx = tx[["Date", "ID_Transaction", "ID_Client", "Agence", "Ville", "Commercial", "ID_Agent", "Produit", "Canal",
             "Montant (FCFA)", "Commission (FCFA)", "Commission_Agent (FCFA)", "Statut", "Motif_Echec"]]
    cl = clients.copy()
    cl.loc[rng.choice(n_clients, 36, replace=False), "Age"] = np.nan
    cl.loc[rng.choice(n_clients, 8, replace=False), "Age"] = [0, 5, 150, 999, 3, 120, 0, 200]
    cl["Sexe"] = cl["Sexe"].astype(object)
    cl.loc[rng.choice(n_clients, 6, replace=False), "Sexe"] = "H"
    cl.loc[rng.choice(n_clients, 15, replace=False), "Sexe"] = np.nan
    ag_out = agents.drop(columns=["_i", "_w"])

    # ───────────── Écriture Excel ─────────────
    chemin.parent.mkdir(parents=True, exist_ok=True)
    with pd.ExcelWriter(chemin, engine="openpyxl", datetime_format="yyyy-mm-dd hh:mm:ss") as w:
        for nom, d in [("Transactions", tx), ("Clients", cl), ("Agents", ag_out), ("Objectifs", objectifs)]:
            d.to_excel(w, sheet_name=nom, index=False)
        _styler(w)

    # ───────────── Feuille Indicateurs : baseline = 1er trimestre, cibles par défaut ─────────────
    from . import cadre, data, kpis
    donnees = data.charger_donnees(chemin)
    base, cibles, _ = kpis.baselines_et_cibles(donnees)
    ind = pd.DataFrame([{"ID_Indicateur": d["id"], "Indicateur": d["nom"], "Unité": d["unite"], "Sens": d["sens"],
                         "Baseline": round(base[d["id"]], 4), "Cible": round(cibles[d["id"]], 4),
                         "Période_Baseline": "1er trimestre des données",
                         "Remarque": "Cible proposée par défaut — à valider avec les parties prenantes"} for d in cadre.DEFS])
    with pd.ExcelWriter(chemin, engine="openpyxl", mode="a", if_sheet_exists="replace") as w:
        ind.to_excel(w, sheet_name="Indicateurs", index=False)
        _styler(w)
    return chemin


def _styler(writer):
    from openpyxl.styles import Alignment, Font, PatternFill
    for ws in writer.book.worksheets:
        for c in ws[1]:
            c.font = Font(name="Arial", bold=True, color="FFFFFF")
            c.fill = PatternFill("solid", fgColor="003366")
            c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.freeze_panes = "A2"
        for i, col in enumerate(ws.iter_cols(min_row=1, max_row=min(ws.max_row, 50)), start=1):
            larg = max(len(str(c.value)) if c.value is not None else 0 for c in col)
            ws.column_dimensions[ws.cell(1, i).column_letter].width = min(max(12, larg + 3), 55)


if __name__ == "__main__":
    ap = argparse.ArgumentParser(description="Génère le classeur de démonstration Mobile Money.")
    ap.add_argument("--sortie", default=str(C.FICHIER_DEFAUT))
    ap.add_argument("--graine", type=int, default=42)
    ap.add_argument("--transactions", type=int, default=50_000)
    a = ap.parse_args()
    print("Fichier généré :", generer(a.sortie, a.graine, a.transactions))
